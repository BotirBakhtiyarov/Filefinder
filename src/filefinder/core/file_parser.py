"""Content extraction from supported file types.

Dispatches :func:`extract_text` by file extension to the appropriate
parser. Heavy/optional parsers (python-docx, openpyxl, python-pptx,
pdfplumber) are imported lazily so this module always imports cleanly.

This module never raises on unreadable content: failures are logged and an
empty string is returned so the indexer can skip the file gracefully.
"""

from __future__ import annotations

import logging
from pathlib import Path

from filefinder.core import ocr

logger = logging.getLogger(__name__)

# Average characters per page below which a PDF is considered scanned and
# OCR fallback is attempted.
_PDF_OCR_THRESHOLD = 20


class UnsupportedFileError(Exception):
    """Raised when a file's extension has no registered text extractor."""


# ----------------------------------------------------------------------
# Per-type extractors (all best-effort; raise only for programmer errors)
# ----------------------------------------------------------------------

def _extract_text_plain(path: Path) -> str:
    """Read a plain-text file trying utf-8, then gbk, then lossy utf-8."""
    raw = path.read_bytes()
    for encoding in ("utf-8", "gbk"):
        try:
            return raw.decode(encoding)
        except (UnicodeDecodeError, LookupError):
            continue
    return raw.decode("utf-8", errors="replace")


def _extract_text_docx(path: Path) -> str:
    """Extract paragraphs and table cell text from a .docx file."""
    from docx import Document  # lazy import (python-docx)

    doc = Document(str(path))
    parts: list[str] = []
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text = cell.text.strip()
                if text:
                    parts.append(text)
    return "\n".join(parts)


def _extract_text_xlsx(path: Path) -> str:
    """Extract all cell values from every sheet of an .xlsx workbook."""
    from openpyxl import load_workbook  # lazy import

    workbook = load_workbook(str(path), read_only=True, data_only=True)
    try:
        parts: list[str] = []
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows(values_only=True):
                values = [str(cell) for cell in row if cell is not None]
                if values:
                    parts.append(" ".join(values))
        return "\n".join(parts)
    finally:
        workbook.close()


def _extract_text_pptx(path: Path) -> str:
    """Extract text from all shapes on all slides of a .pptx file."""
    from pptx import Presentation  # lazy import (python-pptx)

    presentation = Presentation(str(path))
    parts: list[str] = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    parts.append(text)
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        text = cell.text.strip()
                        if text:
                            parts.append(text)
    return "\n".join(parts)


def _extract_text_pdf(path: Path) -> str:
    """Extract text from a PDF; fall back to OCR for scanned documents."""
    import pdfplumber  # lazy import

    parts: list[str] = []
    with pdfplumber.open(str(path)) as pdf:
        page_count = len(pdf.pages)
        for page in pdf.pages:
            try:
                text = page.extract_text() or ""
            except Exception as exc:  # noqa: BLE001 - skip broken page
                logger.warning("Failed to extract text from PDF page (%s): %s", path, exc)
                text = ""
            if text.strip():
                parts.append(text.strip())
    full_text = "\n".join(parts)

    avg_chars = len(full_text) / page_count if page_count else 0
    if avg_chars < _PDF_OCR_THRESHOLD:
        logger.info(
            "PDF %s looks scanned (avg %.1f chars/page); trying OCR fallback.",
            path, avg_chars,
        )
        ocr_text = ocr.extract_text_from_pdf(str(path))
        if ocr_text.strip():
            return ocr_text
    return full_text


def _extract_text_image(path: Path) -> str:
    """Extract text from an image via OCR."""
    return ocr.extract_text_from_image(str(path))


_EXTRACTORS = {
    ".txt": _extract_text_plain,
    ".md": _extract_text_plain,
    ".docx": _extract_text_docx,
    ".xlsx": _extract_text_xlsx,
    ".pptx": _extract_text_pptx,
    ".pdf": _extract_text_pdf,
    ".jpg": _extract_text_image,
    ".jpeg": _extract_text_image,
    ".png": _extract_text_image,
    ".bmp": _extract_text_image,
}


# ----------------------------------------------------------------------
# Public API
# ----------------------------------------------------------------------

def extract_text(path: str) -> str:
    """Extract plain text from a file, dispatching on its extension.

    Args:
        path: filesystem path of the file to parse.

    Returns:
        The extracted text, or ``""`` if the file could not be read or
        contains no extractable text.

    Raises:
        UnsupportedFileError: if the file extension is not supported.
    """
    file_path = Path(path)
    extension = file_path.suffix.lower()
    extractor = _EXTRACTORS.get(extension)
    if extractor is None:
        raise UnsupportedFileError(f"Unsupported file type: {extension or '(none)'}")
    try:
        return extractor(file_path) or ""
    except Exception as exc:  # noqa: BLE001 - never crash on unreadable content
        logger.warning("Failed to extract text from %s: %s", path, exc)
        return ""
